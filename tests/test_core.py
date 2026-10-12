from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from timetrack.database import Database
from timetrack.export import export_project_sheet
from timetrack.formatting import format_hms
from timetrack.timer_service import TimerService


class DatabaseTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "test.db")

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_project_and_track_roundtrip(self) -> None:
        self.db.set_user_name("Wael")
        project = self.db.create_project("Website")
        track = self.db.add_track(
            project.id,
            "Homepage polish",
            "2026-10-12",
            "09:00:00",
            "10:15:30",
            4530,
            "Shipped the new hero section",
        )
        loaded = self.db.get_project(project.id)
        self.assertEqual(loaded.track_count, 1)
        self.assertEqual(loaded.total_seconds, 4530)
        self.assertEqual(track.title, "Homepage polish")
        self.assertEqual(self.db.list_tracks(project.id)[0].output, "Shipped the new hero section")

    def test_unique_project_names(self) -> None:
        self.db.create_project("Design")
        with self.assertRaises(Exception):
            self.db.create_project("design")


class FormatTests(unittest.TestCase):
    def test_hms(self) -> None:
        self.assertEqual(format_hms(0), "00:00:00")
        self.assertEqual(format_hms(3661), "01:01:01")


class TimerTests(unittest.TestCase):
    def test_pause_does_not_run_until_start(self) -> None:
        timer = TimerService()
        snapshot = timer.snapshot()
        self.assertFalse(snapshot.running)
        timer.pause()
        self.assertFalse(timer.is_active())
        timer.shutdown()


class ExportTests(unittest.TestCase):
    def test_excel_headers(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        db = Database(Path(tmp.name) / "test.db")
        db.set_user_name("Wael")
        project = db.create_project("Mobile App")
        db.add_track(project.id, "Login flow", "2026-10-12", "09:00:00", "10:00:00", 3600, "Auth done")
        dest = Path(tmp.name) / "sheet.xlsx"
        export_project_sheet(db, project, "Wael", dest=dest)
        self.assertTrue(dest.exists())
        from openpyxl import load_workbook

        book = load_workbook(dest)
        sheet = book.active
        self.assertEqual(sheet["A3"].value, "User")
        self.assertEqual(sheet["B3"].value, "Wael")
        self.assertEqual(sheet["A4"].value, "Project Title")
        self.assertEqual(sheet["B4"].value, "Mobile App")
        self.assertEqual(sheet["A5"].value, "Total Duration")
        self.assertEqual(sheet["B5"].value, "01:00:00")
        self.assertEqual(sheet["A6"].value, "Number of tracks")
        self.assertEqual(sheet["B6"].value, "1")
        self.assertEqual(
            [sheet.cell(8, col).value for col in range(1, 7)],
            ["Day Date", "Start Time", "End Time", "Total Duration", "Title", "Output"],
        )
        self.assertEqual(sheet["A9"].value, "2026-10-12")
        self.assertEqual(sheet["E9"].value, "Login flow")


if __name__ == "__main__":
    unittest.main()
