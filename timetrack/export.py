from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from timetrack.database import Database, Project
from timetrack.formatting import format_hms
from timetrack.paths import exports_dir


HEADER_FILL = PatternFill("solid", fgColor="1E6FE8")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
LABEL_FONT = Font(name="Calibri", bold=True, color="1E3A5F", size=11)
VALUE_FONT = Font(name="Calibri", color="1E293B", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color="1E6FE8", size=18)
THIN = Border(
    left=Side(style="thin", color="D0D7E2"),
    right=Side(style="thin", color="D0D7E2"),
    top=Side(style="thin", color="D0D7E2"),
    bottom=Side(style="thin", color="D0D7E2"),
)
ZEBRA = PatternFill("solid", fgColor="F4F8FC")


def export_project_sheet(db: Database, project: Project, user_name: str, dest: Path | None = None) -> Path:
    tracks = db.list_tracks(project.id)
    total_seconds = sum(track.duration_seconds for track in tracks)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TrackSheet"

    sheet.merge_cells("A1:F1")
    sheet["A1"] = "TimeTrack — TrackSheet"
    sheet["A1"].font = TITLE_FONT
    sheet["A1"].alignment = Alignment(vertical="center")

    header_rows = [
        ("User", user_name or "—"),
        ("Project Title", project.name),
        ("Total Duration", format_hms(total_seconds)),
        ("Number of tracks", str(len(tracks))),
    ]
    for index, (label, value) in enumerate(header_rows, start=3):
        sheet[f"A{index}"] = label
        sheet[f"A{index}"].font = LABEL_FONT
        sheet.merge_cells(start_row=index, start_column=2, end_row=index, end_column=3)
        sheet[f"B{index}"] = value
        sheet[f"B{index}"].font = VALUE_FONT

    table_start = 8
    columns = [
        "Day Date",
        "Start Time",
        "End Time",
        "Total Duration",
        "Title",
        "Output",
    ]
    for col, heading in enumerate(columns, start=1):
        cell = sheet.cell(table_start, col, heading)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN

    for row_index, track in enumerate(tracks, start=table_start + 1):
        values = [
            track.day_date,
            track.start_time,
            track.end_time,
            format_hms(track.duration_seconds),
            track.title,
            track.output,
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, col, value)
            cell.font = VALUE_FONT
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if (row_index - table_start) % 2 == 0:
                cell.fill = ZEBRA

    widths = [16, 14, 14, 16, 32, 48]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A9"
    sheet.print_title_rows = "1:8"

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    safe_name = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in project.name).strip()
    path = dest or (exports_dir() / f"{safe_name}_TrackSheet_{stamp}.xlsx")
    workbook.save(path)
    return Path(path)
